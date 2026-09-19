"""run_sumslice_llm_statistics.py

Significance tests for the four LLMs on the SumSlice sample.

  1. Friedman + Kendall's W per judge criterion (method-level points averaged
     over presentations), Holm-corrected Wilcoxon post-hoc where significant.
  2. The same on the automated metrics (cosine, BERTScore F1).
  3. Both, per project (25 methods each) as a secondary analysis.
  4. Spearman correlation between automated metrics and judge points.

Inputs:  evaluation_results/judge_rankings_long.csv, metrics_per_item.csv
Outputs: evaluation_results/SUMSLICE_LLM_Statistical_Results.md / .xlsx

:author: Najam Nazar
:version: 2.0.0
:license: MIT
"""

from __future__ import annotations

import sys

import pandas as pd
from openpyxl import Workbook

from sumslice_common import DISPLAY_NAMES, REPO_ROOT, RESULTS_DIR

sys.path.insert(0, str(REPO_ROOT / "resources"))
from stats_shared import (  # noqa: E402
    add_heatmap, run_friedman_and_posthoc, run_spearman_sensitivity, write_df, write_pivot_block,
)

AUTO = {"cosine_similarity": "Cosine (TF-IDF)", "bert_f1": "BERTScore F1"}
FLOATS = ["Friedman chi2", "p-value", "Kendall's W", "Mean A", "Mean B", "Wilcoxon stat",
          "p-value (raw)", "p-value (Holm)", "Spearman rho"]


def md(df: pd.DataFrame, cols) -> list:
    if df.empty:
        return ["_none_", ""]
    lines = ["| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
    for _, r in df.iterrows():
        lines.append("| " + " | ".join(
            f"{r[c]:.4g}" if isinstance(r[c], float) else ("Yes" if r[c] is True else "No" if r[c] is False else str(r[c]))
            for c in cols) + " |")
    return lines + [""]


def main() -> None:
    jp, mp = RESULTS_DIR / "judge_rankings_long.csv", RESULTS_DIR / "metrics_per_item.csv"
    for p in (jp, mp):
        if not p.exists():
            raise SystemExit(f"Missing {p}; run evaluate_sumslice_summaries.py and rank_sumslice_summaries.py first.")
    judge = pd.read_csv(jp)
    judge = judge.groupby(["project", "method_id", "criterion", "system"])["points"].mean().reset_index()
    judge["item"] = judge["project"] + "::" + judge["method_id"].astype(str)
    judge["model"] = judge["system"].map(DISPLAY_NAMES)
    metrics = pd.read_csv(mp)
    metrics["item"] = metrics["project"] + "::" + metrics["method_id"].astype(str)
    metrics["model"] = metrics["system"].map(DISPLAY_NAMES)
    auto = metrics.melt(id_vars=["item", "project", "model"],
                        value_vars=[c for c in AUTO if metrics[c].notna().any()],
                        var_name="criterion", value_name="points")
    auto["criterion"] = auto["criterion"].map(AUTO)

    frs, phs = [], []
    for name, frame in (("judge", judge), ("automated", auto)):
        for scope, g in [("all", frame)] + [(f"project={p}", gp) for p, gp in frame.groupby("project")]:
            fr, ph = run_friedman_and_posthoc(g, criterion_col="criterion", item_col="item",
                                              model_col="model", value_col="points")
            for d in (fr, ph):
                if not d.empty:
                    d.insert(0, "Scope", f"{name} / {scope}")
            frs.append(fr)
            phs.append(ph)
    friedman = pd.concat(frs, ignore_index=True)
    posthoc = pd.concat(phs, ignore_index=True)

    sens = [run_spearman_sensitivity(judge.rename(columns={"points": "judge_points"}), metrics,
                                     item_col="item", model_col="model", criterion_col="criterion",
                                     rubric_value_col="judge_points", metric_col=c, metric_name=l)
            for c, l in AUTO.items() if metrics[c].notna().any()]
    sens = pd.concat(sens, ignore_index=True) if sens else pd.DataFrame()
    means = judge.groupby(["model", "criterion"])["points"].mean().unstack()

    lines = ["# SUMSLICE_LLM statistical results", "",
             "Reference: SumSlice tool summaries. Judge points: 1st=4 ... 4th=1, averaged over "
             "presentation orders per method. Project-level tests use 25 methods each.", "",
             "## Mean judge points", "", means.round(3).to_markdown(), "", "## Friedman tests", ""]
    lines += md(friedman, ["Scope", "Criterion", "N Items", "K Models", "Friedman chi2", "p-value",
                           "Kendall's W", "Significant (a=.05)"])
    lines += ["## Holm-corrected Wilcoxon post-hoc", ""]
    lines += md(posthoc, ["Scope", "Criterion", "Model A", "Model B", "Mean A", "Mean B",
                          "p-value (Holm)", "Significant (a=.05)"])
    if not sens.empty:
        for name, g in sens.groupby("Automated Metric"):
            lines += [f"## Spearman: {name} vs judge points", "",
                      g.pivot(index="Model", columns="Criterion", values="Spearman rho").round(3).to_markdown(), ""]
    (RESULTS_DIR / "SUMSLICE_LLM_Statistical_Results.md").write_text("\n".join(lines))

    wb = Workbook()
    wb.remove(wb.active)
    write_pivot_block(wb.create_sheet("Mean judge points"), means, 1, 1, "Mean judge points")
    write_df(wb.create_sheet("Friedman"), friedman, sig_col="Significant (a=.05)", float_cols=FLOATS)
    ws = wb.create_sheet("Wilcoxon post-hoc")
    if posthoc.empty:
        ws.cell(row=1, column=1, value="No Friedman test reached significance.")
    else:
        write_df(ws, posthoc, sig_col="Significant (a=.05)", float_cols=FLOATS)
    for name, g in (sens.groupby("Automated Metric") if not sens.empty else []):
        ws = wb.create_sheet(f"Spearman {name}"[:31])
        g = g.reset_index(drop=True)
        write_df(ws, g, sig_col="Significant (a=.05)", float_cols=FLOATS)
        pv = g.pivot(index="Model", columns="Criterion", values="Spearman rho")
        _, r0, r1, c0, c1 = write_pivot_block(ws, pv, len(g) + 4, 1, "Spearman rho")
        add_heatmap(ws, r0, r1, c0 + 1, c1)
    wb.save(RESULTS_DIR / "SUMSLICE_LLM_Statistical_Results.xlsx")
    print(means.round(3).to_string())
    print(f"\nFriedman tests: {len(friedman)}; wrote {RESULTS_DIR}")


if __name__ == "__main__":
    main()
