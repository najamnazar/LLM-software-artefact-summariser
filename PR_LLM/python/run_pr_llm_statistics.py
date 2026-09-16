"""Statistical significance analysis for PR-LLM, at the same rigor as DPS-LLM's
section 3.7 (Friedman + Kendall's W + Holm-corrected pairwise Wilcoxon +
Spearman sensitivity).

The comparison set mirrors DPS-LLM: the deterministic tool baseline is tested
alongside the LLMs, all judged against the human gold descriptions.

Inputs:
  PR_LLM/output/rubric_eval_checkpoint.json
      Per-item, per-criterion 1..k place RANKS (1 = best) written by
      rank_pr_summaries.py, keyed by pr_id. Converted here to
      POINTS = (k + 1) - rank so 1st place scores k points and last scores 1.
  PR_LLM/evaluation-results/<SYSTEM>_vs_human_scores.jsonl
      Per-item automated metrics (bert_precision/recall/f1, cosine_similarity)
      for each system, written by evaluate_pr_summaries.py and keyed by pr_id.

Analyses:
  1. Friedman omnibus test per criterion (items x systems) + Kendall's W.
  2. Holm-corrected pairwise Wilcoxon signed-rank post-hoc where Friedman is significant.
  3. Spearman sensitivity: automated metric (BERT F1, Cosine Similarity) vs.
     judge-rubric points, per system per criterion.

Output: PR_LLM_Statistical_Results.xlsx and PR_LLM_Statistical_Results.md
in PR_LLM/output/
"""

import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent / "resources"))

import pr_corpus
from pr_corpus import DISPLAY_NAMES, OUTPUT_DIR, RESULTS_DIR
from stats_shared import (
    run_friedman_and_posthoc, run_spearman_sensitivity,
    write_df, write_pivot_block, add_grouped_bar_chart, add_heatmap,
)

CHECKPOINT = OUTPUT_DIR / "rubric_eval_checkpoint.json"


def load_checkpoint_payload():
    """Read the rubric checkpoint, returning (entries, systems).

    Accepts the current ``{"meta": ..., "entries": ...}`` layout and the older
    flat ``{item_id: record}`` layout, so a checkpoint written before the
    corpus change still loads.
    """
    if not CHECKPOINT.exists():
        raise FileNotFoundError(
            f"Rubric checkpoint not found: {CHECKPOINT}\n"
            "Run rank_pr_summaries.py first."
        )

    data = json.loads(CHECKPOINT.read_text())
    if isinstance(data, dict) and "entries" in data:
        entries = data["entries"]
        systems = (data.get("meta") or {}).get("systems")
    else:
        entries = data
        systems = None

    if not systems:
        # Derive the system set from the ranks themselves when meta is absent.
        found: list[str] = []
        for item in entries.values():
            for scores in (item.get("criteria") or {}).values():
                for key in (scores or {}):
                    if key not in found:
                        found.append(key)
        systems = found

    if not systems:
        raise ValueError(f"No systems found in {CHECKPOINT}")
    return entries, systems


def load_rubric_long(entries, systems):
    """Flatten the checkpoint into one row per (item, criterion, system)."""
    k = len(systems)
    rows = []
    for key, item in entries.items():
        item_id = item.get("pr_id") or item.get("id") or key
        for criterion, scores in (item.get("criteria") or {}).items():
            if not scores:
                continue
            for system, rank in scores.items():
                if rank is None:
                    continue
                rows.append({
                    "item_id": item_id,
                    "criterion": criterion,
                    "model": DISPLAY_NAMES.get(system, system),
                    "rank": rank,
                    "points": (k + 1) - rank,
                })
    return pd.DataFrame(rows)


def load_metrics_long(systems):
    """Load per-item automated metrics for each system from evaluation-results/."""
    rows = []
    for system in systems:
        path = RESULTS_DIR / f"{system}_vs_human_scores.jsonl"
        if not path.exists():
            print(f"  WARNING: {path.name} not found — skipping {system} metrics")
            continue
        for rec in pr_corpus.load_jsonl(path):
            metrics = rec.get("metrics", {})
            rows.append({
                "item_id": rec.get("pr_id"),
                "model": DISPLAY_NAMES.get(system, system),
                "bert_f1": metrics.get("bert_f1"),
                "cosine_similarity": metrics.get("cosine_similarity"),
            })
    return pd.DataFrame(rows)


def write_markdown(friedman_df, posthoc_df, sens_df, n_systems):
    lines = []
    lines.append("# PR-LLM Statistical Significance Tests\n")
    lines.append("Mirrors DPS-LLM's ALL_RESULTS.md section 3.7. The deterministic tool baseline is "
                 "tested alongside the LLMs. Ratings are 1st-to-last place judge rankings converted "
                 f"to points (1st={n_systems} ... last=1); non-parametric tests used throughout "
                 "since the data is ordinal.\n")

    lines.append(f"## Friedman Omnibus Tests (per criterion, k={n_systems} systems)\n")
    if friedman_df.empty:
        lines.append("_No criterion had enough complete items to test._\n")
    else:
        lines.append("| Criterion | N Items | K Models | Friedman chi2 | p-value | Kendall's W | Significant |")
        lines.append("|---|---|---|---|---|---|---|")
        for _, r in friedman_df.iterrows():
            kendalls_w = r["Kendall's W"]
            lines.append(f"| {r['Criterion']} | {r['N Items']} | {r['K Models']} | {r['Friedman chi2']:.4f} | "
                         f"{r['p-value']:.4g} | {kendalls_w:.4f} | "
                         f"{'Yes' if r['Significant (a=.05)'] else 'No'} |")
        lines.append("")

    lines.append("## Holm-Corrected Pairwise Wilcoxon Post-Hoc (only for significant Friedman criteria)\n")
    if posthoc_df.empty:
        lines.append("_No criterion reached Friedman significance; no post-hoc tests run._\n")
    else:
        lines.append("| Criterion | Model A | Model B | Mean A | Mean B | p (raw) | p (Holm) | Significant |")
        lines.append("|---|---|---|---|---|---|---|---|")
        for _, r in posthoc_df.iterrows():
            lines.append(f"| {r['Criterion']} | {r['Model A']} | {r['Model B']} | {r['Mean A']:.3f} | "
                         f"{r['Mean B']:.3f} | {r['p-value (raw)']:.4g} | {r['p-value (Holm)']:.4g} | "
                         f"{'Yes' if r['Significant (a=.05)'] else 'No'} |")
        lines.append("")

    lines.append("## Spearman Sensitivity: Automated Metrics vs. Judge-Rubric Points\n")
    if sens_df.empty:
        lines.append("_No overlapping rubric and metric rows to correlate._\n")
    else:
        for metric_name in sens_df["Automated Metric"].unique():
            lines.append(f"### vs. {metric_name}\n")
            sub = sens_df[sens_df["Automated Metric"] == metric_name]
            pivot = sub.pivot(index="Model", columns="Criterion", values="Spearman rho").round(4)
            lines.append("| Model | " + " | ".join(pivot.columns) + " |")
            lines.append("|---|" + "---|" * len(pivot.columns))
            for model, row in pivot.iterrows():
                lines.append(f"| {model} | " + " | ".join(f"{v:.4f}" if pd.notna(v) else "n/a" for v in row) + " |")
            lines.append("")

    (OUTPUT_DIR / "PR_LLM_Statistical_Results.md").write_text("\n".join(lines))


def write_excel(friedman_df, posthoc_df, sens_df):
    float_cols = ["Friedman chi2", "p-value", "Kendall's W", "Mean A", "Mean B",
                  "Wilcoxon stat", "p-value (raw)", "p-value (Holm)", "Spearman rho"]

    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Friedman (omnibus)")
    if not friedman_df.empty:
        write_df(ws, friedman_df, sig_col="Significant (a=.05)", float_cols=float_cols)
        pivot = friedman_df.set_index("Criterion")[["Kendall's W"]]
        pivot.index.name = "Criterion"
        header_row, first_data_row, last_data_row, first_col, last_col = write_pivot_block(
            ws, pivot, len(friedman_df) + 4, 1, "Effect size by criterion")
        add_grouped_bar_chart(ws, header_row, first_data_row, last_data_row, first_col, first_col + 1, last_col,
                              "PR-LLM: Kendall's W by criterion", "Kendall's W", "Criterion",
                              f"{get_column_letter(last_col + 2)}{header_row}")
    else:
        ws.cell(row=1, column=1, value="No criterion had enough complete items to test.")

    ws = wb.create_sheet("Wilcoxon post-hoc")
    if not posthoc_df.empty:
        write_df(ws, posthoc_df, sig_col="Significant (a=.05)", float_cols=float_cols)
    else:
        ws.cell(row=1, column=1, value="No omnibus Friedman test reached significance; no post-hoc tests run.")

    if sens_df.empty:
        ws = wb.create_sheet("Spearman")
        ws.cell(row=1, column=1, value="No overlapping rubric and metric rows to correlate.")
    else:
        for metric_name in sens_df["Automated Metric"].unique():
            ws = wb.create_sheet(f"Spearman vs {metric_name}"[:31])
            sub = sens_df[sens_df["Automated Metric"] == metric_name].reset_index(drop=True)
            write_df(ws, sub, sig_col="Significant (a=.05)", float_cols=float_cols)
            pivot = sub.pivot(index="Model", columns="Criterion", values="Spearman rho")
            header_row, first_data_row, last_data_row, first_col, last_col = write_pivot_block(
                ws, pivot, len(sub) + 4, 1, f"Spearman rho vs {metric_name}")
            add_heatmap(ws, first_data_row, last_data_row, first_col + 1, last_col)

    wb.save(OUTPUT_DIR / "PR_LLM_Statistical_Results.xlsx")


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    entries, systems = load_checkpoint_payload()
    print(f"Systems compared: {', '.join(DISPLAY_NAMES.get(s, s) for s in systems)}")

    rubric_long = load_rubric_long(entries, systems)
    if rubric_long.empty:
        raise SystemExit("No rubric rankings found in the checkpoint — nothing to test.")

    metrics_long = load_metrics_long(systems)

    friedman_df, posthoc_df = run_friedman_and_posthoc(
        rubric_long, criterion_col="criterion", item_col="item_id", model_col="model", value_col="points")

    if metrics_long.empty:
        sens_df = pd.DataFrame(columns=["Model", "Criterion", "Automated Metric",
                                        "N", "Spearman rho", "p-value", "Significant (a=.05)"])
    else:
        sens_bert = run_spearman_sensitivity(
            rubric_long, metrics_long, item_col="item_id", model_col="model", criterion_col="criterion",
            rubric_value_col="points", metric_col="bert_f1", metric_name="BERT F1")
        sens_cos = run_spearman_sensitivity(
            rubric_long, metrics_long, item_col="item_id", model_col="model", criterion_col="criterion",
            rubric_value_col="points", metric_col="cosine_similarity", metric_name="Cosine Similarity")
        sens_df = pd.concat([sens_bert, sens_cos], ignore_index=True)

    write_markdown(friedman_df, posthoc_df, sens_df, len(systems))
    write_excel(friedman_df, posthoc_df, sens_df)

    n_sig = int(friedman_df["Significant (a=.05)"].sum()) if not friedman_df.empty else 0
    print(f"Friedman: {len(friedman_df)} criteria tested, {n_sig} significant")
    print(f"Wilcoxon post-hoc: {len(posthoc_df)} pairwise comparisons")
    print(f"Spearman sensitivity: {len(sens_df)} (model, criterion, metric) rows")
    print(f"Wrote results to {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
